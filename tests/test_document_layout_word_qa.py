import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from scripts.document_layout_word_qa import inspect_docx


class WordLayoutQATest(unittest.TestCase):
    def make_doc(self, inconsistent=False):
        temp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        temp.close()
        path = Path(temp.name)
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        table.style = "Table Grid"
        values = [
            ["Item", "Amount (EUR)", "Explanation"],
            ["Rollout", "1200", "Production cost"],
            ["Props", "300", "One-time tooling"],
        ]
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                cell.text = values[ri][ci]
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if ci == 1 else WD_ALIGN_PARAGRAPH.LEFT
        if inconsistent:
            table.cell(2, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        doc.save(path)
        return path

    def test_consistent_docx_has_no_major_alignment_issues(self):
        path = self.make_doc()
        try:
            report = inspect_docx(path)
            self.assertEqual(report["tables"], 1)
            self.assertFalse(any(i["severity"] == "major" for i in report["issues"]))
        finally:
            path.unlink(missing_ok=True)

    def test_flags_misaligned_numeric_column(self):
        path = self.make_doc(inconsistent=True)
        try:
            report = inspect_docx(path)
            codes = {issue["code"] for issue in report["issues"]}
            self.assertIn("MIXED_COLUMN_ALIGNMENT", codes)
            self.assertIn("NUMERIC_CELL_LEFT_ALIGNED", codes)
        finally:
            path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
