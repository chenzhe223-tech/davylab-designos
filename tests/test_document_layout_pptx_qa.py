import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from scripts.document_layout_pptx_qa import inspect_pptx


class PowerPointLayoutQATest(unittest.TestCase):
    def make_presentation(self, inconsistent=False):
        temp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        temp.close()
        path = Path(temp.name)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table_shape = slide.shapes.add_table(3, 3, 500000, 500000, 8000000, 2500000)
        table = table_shape.table
        values = [
            ["Item", "Amount (EUR)", "Explanation"],
            ["Rollout", "1200", "Production cost"],
            ["Props", "300", "One-time tooling"],
        ]
        for ri in range(3):
            for ci in range(3):
                cell = table.cell(ri, ci)
                cell.text = values[ri][ci]
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment = PP_ALIGN.RIGHT if ci == 1 else PP_ALIGN.LEFT
        if inconsistent:
            table.cell(2, 1).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        prs.save(path)
        return path

    def test_consistent_presentation_has_no_major_issues(self):
        path = self.make_presentation()
        try:
            report = inspect_pptx(path)
            self.assertEqual(report["slides"], 1)
            self.assertFalse(any(i["severity"] == "major" for i in report["issues"]))
        finally:
            path.unlink(missing_ok=True)

    def test_flags_misaligned_numeric_table_column(self):
        path = self.make_presentation(inconsistent=True)
        try:
            report = inspect_pptx(path)
            codes = {issue["code"] for issue in report["issues"]}
            self.assertIn("MIXED_COLUMN_ALIGNMENT", codes)
            self.assertIn("NUMERIC_CELL_LEFT_ALIGNED", codes)
        finally:
            path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
