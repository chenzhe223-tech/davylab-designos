import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from scripts.document_layout_excel_qa import inspect_workbook


class ExcelLayoutQATest(unittest.TestCase):
    def make_book(self, inconsistent=False):
        temp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp.close()
        path = Path(temp.name)
        wb = Workbook()
        ws = wb.active
        ws.title = "Cost summary"
        ws.append(["Item", "Amount (EUR)", "Explanation"])
        ws.append(["Rollout", 1200.0, "Production cost"])
        ws.append(["Props", 300.0, "One-time tooling"])
        for row in ws.iter_rows(min_row=1, max_row=3):
            row[0].alignment = Alignment(horizontal="left", vertical="top")
            row[1].alignment = Alignment(horizontal="right", vertical="top")
            row[2].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws["B2"].number_format = '€#,##0.00'
        ws["B3"].number_format = '€#,##0.00'
        if inconsistent:
            ws["B3"].alignment = Alignment(horizontal="left")
            ws["B3"].number_format = '0.00'
        wb.save(path)
        wb.close()
        return path

    def test_consistent_workbook_has_no_major_issues(self):
        path = self.make_book()
        try:
            report = inspect_workbook(path)
            self.assertEqual(report["sheets"], 1)
            self.assertFalse(any(i["severity"] == "major" for i in report["issues"]))
        finally:
            path.unlink(missing_ok=True)

    def test_flags_alignment_and_number_format_inconsistency(self):
        path = self.make_book(inconsistent=True)
        try:
            report = inspect_workbook(path)
            codes = {issue["code"] for issue in report["issues"]}
            self.assertIn("MIXED_EXPLICIT_COLUMN_ALIGNMENT", codes)
            self.assertIn("NUMERIC_CELL_LEFT_ALIGNED", codes)
            self.assertIn("MIXED_NUMERIC_FORMATS", codes)
        finally:
            path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
