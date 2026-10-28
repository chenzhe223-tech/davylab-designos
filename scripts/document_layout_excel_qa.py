#!/usr/bin/env python3
"""Static Excel layout QA. Reports likely style inconsistencies; does not auto-repair workbooks."""
import argparse
import json
import os
import sys
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Missing dependency: install requirements-layout-qa.txt (openpyxl).", file=sys.stderr)
    raise SystemExit(2)


def inspect_workbook(filename):
    path = Path(filename)
    workbook = load_workbook(path, data_only=False, read_only=False)
    issues = []

    def add(severity, code, message, sheet=None, cell=None):
        issue = {"severity": severity, "code": code, "message": message}
        if sheet is not None:
            issue["sheet"] = sheet
        if cell is not None:
            issue["cell"] = cell
        issues.append(issue)

    for ws in workbook.worksheets:
        columns = defaultdict(list)
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                value = cell.value
                alignment = cell.alignment.horizontal
                numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
                formula = isinstance(value, str) and value.startswith("=")
                if alignment:
                    columns[cell.column].append((cell, alignment, numeric or formula, value))
                if numeric and alignment == "left":
                    add("warning", "NUMERIC_CELL_LEFT_ALIGNED",
                        f"Numeric cell {cell.coordinate} is explicitly left-aligned.", ws.title, cell.coordinate)
                if isinstance(value, str) and "\n" in value and not cell.alignment.wrap_text:
                    add("info", "MULTILINE_WITHOUT_WRAP",
                        f"Cell {cell.coordinate} contains line breaks but wrap_text is disabled.", ws.title, cell.coordinate)
        for col_idx, entries in columns.items():
            explicit = [(cell, align, is_numeric, value) for cell, align, is_numeric, value in entries if align not in ("general", None)]
            aligns = sorted({align for _, align, _, _ in explicit})
            if len(aligns) > 1:
                col_letter = ws.cell(row=1, column=col_idx).column_letter
                add("major", "MIXED_EXPLICIT_COLUMN_ALIGNMENT",
                    f"Column {col_letter} has conflicting explicit horizontal alignments: {', '.join(aligns)}.", ws.title)
            numeric_formats = {cell.number_format for cell, _, is_numeric, _ in entries if is_numeric}
            if len(numeric_formats) > 1:
                col_letter = ws.cell(row=1, column=col_idx).column_letter
                add("warning", "MIXED_NUMERIC_FORMATS",
                    f"Column {col_letter} uses multiple number formats for numeric/formula cells: {', '.join(sorted(numeric_formats))}.", ws.title)
        for merged in ws.merged_cells.ranges:
            if merged.min_col != merged.max_col and merged.min_row != merged.max_row:
                add("info", "LARGE_MERGED_RANGE",
                    f"Merged range {merged} spans multiple rows and columns; verify it does not interfere with sorting or table navigation.", ws.title)
        if ws.sheet_view.showGridLines is False and not any(cell.has_style for row in ws.iter_rows() for cell in row):
            add("info", "GRIDLINES_HIDDEN_NO_CELL_STYLES",
                "Gridlines are hidden but no styled cells were detected; verify table boundaries remain visible.", ws.title)

    workbook.close()
    summary = {key: 0 for key in ("blocker", "major", "warning", "info")}
    for issue in issues:
        summary[issue["severity"]] += 1
    return {
        "tool": "DAVYLAB Document Layout Excel QA",
        "version": "0.1.0",
        "mode": "static-workbook-style-analysis",
        "file": str(path),
        "sheets": len(load_workbook(path, read_only=True, data_only=True).sheetnames),
        "summary": summary,
        "issues": issues,
        "limitations": [
            "This is a static style heuristic, not a rendered visual inspection.",
            "It does not verify formulas, business correctness, print pagination, or actual cell clipping."
        ]
    }


def main():
    parser = argparse.ArgumentParser(description="Static Excel workbook layout QA")
    parser.add_argument("file", help="Path to .xlsx or .xlsm workbook")
    parser.add_argument("--json", action="store_true", help="Emit JSON report")
    args = parser.parse_args()
    if not os.path.isfile(args.file):
        print(f"Input file not found: {args.file}", file=sys.stderr)
        return 2
    try:
        report = inspect_workbook(args.file)
    except Exception as exc:
        print(f"Unable to inspect workbook: {exc}", file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(f"Excel Layout QA — {report['file']}")
        print(f"Sheets: {report['sheets']}; issues: {len(report['issues'])}")
        for issue in report["issues"]:
            print(f"[{issue['severity'].upper()}] {issue['code']}: {issue['message']}")
        print("Limitations: static analysis only; inspect the rendered workbook/PDF for final visual acceptance.")
    return 1 if any(i["severity"] in ("blocker", "major") for i in report["issues"]) else 0


if __name__ == "__main__":
    raise SystemExit(main())
