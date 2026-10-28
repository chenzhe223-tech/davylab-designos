#!/usr/bin/env python3
"""Static Word DOCX table layout QA. Does not render or modify the document."""
import argparse
import json
import os
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Missing dependency: install requirements-layout-qa.txt (python-docx).", file=sys.stderr)
    raise SystemExit(2)


ALIGNMENTS = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}


def inspect_docx(filename):
    path = Path(filename)
    doc = Document(path)
    issues = []

    def add(severity, code, message, table=None, cell=None):
        item = {"severity": severity, "code": code, "message": message}
        if table is not None:
            item["table"] = table
        if cell is not None:
            item["cell"] = cell
        issues.append(item)

    for table_index, table in enumerate(doc.tables, start=1):
        row_widths = [len(row.cells) for row in table.rows]
        if len(set(row_widths)) > 1:
            add("warning", "INCONSISTENT_ROW_CELL_COUNT",
                "Table has rows with different cell counts; verify intentional merged cells.", table_index)
        max_cols = max(row_widths, default=0)
        for col_idx in range(max_cols):
            cell_alignments = set()
            vertical_alignments = set()
            widths = set()
            for row_idx, row in enumerate(table.rows, start=1):
                if col_idx >= len(row.cells):
                    continue
                cell = row.cells[col_idx]
                widths.add(cell.width)
                if cell.vertical_alignment is not None:
                    vertical_alignments.add(str(cell.vertical_alignment))
                for paragraph in cell.paragraphs:
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    alignment = ALIGNMENTS.get(paragraph.alignment)
                    if alignment:
                        cell_alignments.add(alignment)
                    is_numeric = False
                    try:
                        cleaned = text.replace("€", "").replace("$", "").replace("£", "").replace(",", "").replace(".", "").replace("%", "").replace(" ", "")
                        is_numeric = bool(cleaned) and cleaned.lstrip("+-").isdigit()
                    except Exception:
                        pass
                    if is_numeric and alignment == "left":
                        add("warning", "NUMERIC_CELL_LEFT_ALIGNED",
                            f"Numeric-looking cell at row {row_idx}, column {col_idx + 1} is left-aligned.",
                            table_index, f"R{row_idx}C{col_idx + 1}")
            if len(cell_alignments) > 1:
                add("major", "MIXED_COLUMN_ALIGNMENT",
                    f"Column {col_idx + 1} has conflicting paragraph alignments: {', '.join(sorted(cell_alignments))}.",
                    table_index)
            if len(vertical_alignments) > 1:
                add("warning", "MIXED_COLUMN_VERTICAL_ALIGNMENT",
                    f"Column {col_idx + 1} uses different vertical cell alignments.", table_index)
            nonempty_widths = {w for w in widths if w is not None}
            if len(nonempty_widths) > 1:
                add("info", "MIXED_COLUMN_WIDTHS",
                    f"Column {col_idx + 1} has varying cell width metadata; verify the table grid and merged cells.", table_index)

        if table.style is None:
            add("info", "TABLE_WITHOUT_STYLE",
                "Table has no named Word table style; verify formatting is intentional and reusable.", table_index)

    summary = {key: 0 for key in ("blocker", "major", "warning", "info")}
    for issue in issues:
        summary[issue["severity"]] += 1
    return {
        "tool": "DAVYLAB Document Layout Word QA",
        "version": "0.1.0",
        "mode": "static-docx-style-analysis",
        "file": str(path),
        "tables": len(doc.tables),
        "paragraphs": len(doc.paragraphs),
        "summary": summary,
        "issues": issues,
        "limitations": [
            "This is a static DOCX style heuristic, not a rendered visual inspection.",
            "It does not guarantee no clipping, verify page breaks, or evaluate business calculations."
        ]
    }


def main():
    parser = argparse.ArgumentParser(description="Static Word DOCX layout QA")
    parser.add_argument("file", help="Path to .docx file")
    parser.add_argument("--json", action="store_true", help="Emit JSON report")
    args = parser.parse_args()
    if not os.path.isfile(args.file):
        print(f"Input file not found: {args.file}", file=sys.stderr)
        return 2
    try:
        report = inspect_docx(args.file)
    except Exception as exc:
        print(f"Unable to inspect DOCX: {exc}", file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(f"Word Layout QA — {report['file']}")
        print(f"Tables: {report['tables']}; paragraphs: {report['paragraphs']}; issues: {len(report['issues'])}")
        for issue in report["issues"]:
            print(f"[{issue['severity'].upper()}] {issue['code']}: {issue['message']}")
        print("Limitations: static analysis only; render DOCX to PDF/images for final visual acceptance.")
    return 1 if any(i["severity"] in ("blocker", "major") for i in report["issues"]) else 0


if __name__ == "__main__":
    raise SystemExit(main())
