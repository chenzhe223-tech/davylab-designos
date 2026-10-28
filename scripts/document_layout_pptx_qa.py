#!/usr/bin/env python3
"""Static PowerPoint PPTX layout QA. Reports table alignment and out-of-slide geometry risks."""
import argparse
import json
import os
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("Missing dependency: install requirements-layout-qa.txt (python-pptx).", file=sys.stderr)
    raise SystemExit(2)


ALIGNMENTS = {
    PP_ALIGN.LEFT: "left",
    PP_ALIGN.CENTER: "center",
    PP_ALIGN.RIGHT: "right",
    PP_ALIGN.JUSTIFY: "justify",
}


def numeric_text(text):
    cleaned = (text or "").strip().replace("€", "").replace("$", "").replace("£", "").replace(",", "").replace(".", "").replace("%", "").replace(" ", "")
    return bool(cleaned) and cleaned.lstrip("+-").isdigit()


def inspect_pptx(filename):
    path = Path(filename)
    presentation = Presentation(path)
    issues = []

    def add(severity, code, message, slide=None, shape=None):
        item = {"severity": severity, "code": code, "message": message}
        if slide is not None:
            item["slide"] = slide
        if shape is not None:
            item["shape"] = shape
        issues.append(item)

    slide_width, slide_height = presentation.slide_width, presentation.slide_height
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            label = shape.name or f"Shape {shape_index}"
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > slide_width or shape.top + shape.height > slide_height:
                add("major", "SHAPE_OUTSIDE_SLIDE",
                    f"{label} extends beyond the slide canvas; inspect clipping and margins.", slide_index, label)
            if not getattr(shape, "has_table", False):
                continue
            table = shape.table
            max_cols = len(table.columns)
            for col_idx in range(max_cols):
                aligns = set()
                for row_idx, row in enumerate(table.rows, start=1):
                    cell = row.cells[col_idx]
                    for paragraph in cell.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if not text:
                            continue
                        alignment = ALIGNMENTS.get(paragraph.alignment)
                        if alignment:
                            aligns.add(alignment)
                        if numeric_text(text) and alignment == "left":
                            add("warning", "NUMERIC_CELL_LEFT_ALIGNED",
                                f"Table cell R{row_idx}C{col_idx + 1} contains numeric-looking text and is left-aligned.",
                                slide_index, label)
                if len(aligns) > 1:
                    add("major", "MIXED_COLUMN_ALIGNMENT",
                        f"Table column {col_idx + 1} has conflicting paragraph alignments: {', '.join(sorted(aligns))}.",
                        slide_index, label)
    summary = {key: 0 for key in ("blocker", "major", "warning", "info")}
    for issue in issues:
        summary[issue["severity"]] += 1
    return {
        "tool": "DAVYLAB Document Layout PowerPoint QA",
        "version": "0.1.0",
        "mode": "static-pptx-geometry-and-table-analysis",
        "file": str(path),
        "slides": len(presentation.slides),
        "summary": summary,
        "issues": issues,
        "limitations": [
            "Static geometry and table-style analysis only; does not render slides or detect all text overflow.",
            "Final acceptance requires rendering slides to images and inspecting them at presentation size."
        ]
    }


def main():
    parser = argparse.ArgumentParser(description="Static PowerPoint PPTX layout QA")
    parser.add_argument("file", help="Path to .pptx file")
    parser.add_argument("--json", action="store_true", help="Emit JSON report")
    args = parser.parse_args()
    if not os.path.isfile(args.file):
        print(f"Input file not found: {args.file}", file=sys.stderr)
        return 2
    try:
        report = inspect_pptx(args.file)
    except Exception as exc:
        print(f"Unable to inspect PPTX: {exc}", file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(f"PowerPoint Layout QA — {report['file']}")
        print(f"Slides: {report['slides']}; issues: {len(report['issues'])}")
        for issue in report["issues"]:
            print(f"[{issue['severity'].upper()}] {issue['code']}: {issue['message']}")
        print("Limitations: static analysis only; render slides to images for final visual acceptance.")
    return 1 if any(i["severity"] in ("blocker", "major") for i in report["issues"]) else 0


if __name__ == "__main__":
    raise SystemExit(main())
